from contextlib import contextmanager
import hashlib
import time
import zipfile
import pytest
from sqlalchemy import func,select,text
from app.core.config import get_settings
from app.db.models import DataImport,EventsJournal,ObjectCatalogue,ChannelCatalogue
from app.db.platform import AuditEntry
from app.db.session import SessionLocal,engine
from app.ingestion import limits,import_registry,formats
from test_workflows import file_content


def test_global_slots_reject_immediately_without_registry_then_retry(login):
    c=login();settings=get_settings()
    with engine.connect() as owner:
        for slot in range(settings.import_max_concurrent):
            owner.execute(text('SELECT pg_advisory_lock(:namespace,:slot)'),{'namespace':limits.IMPORT_NAMESPACE,'slot':slot})
        owner.commit()
        try:
            response=c.post('/api/v1/imports/events',files={'file':('busy.csv',file_content('csv'))})
            assert response.status_code==429
            assert response.json()['detail']['code']=='IMPORT_CAPACITY_BUSY'
            assert int(response.headers['retry-after'])>=1
            assert c.get('/api/v1/map/objects').status_code==200
            with SessionLocal() as db:
                assert db.scalar(select(func.count()).select_from(DataImport))==0
                audit=db.scalar(select(AuditEntry).where(AuditEntry.target=='/api/v1/imports/events',AuditEntry.outcome==429))
                assert audit and audit.actor=='dispatcher' and audit.role=='dispatcher'
        finally:
            for slot in range(settings.import_max_concurrent):owner.execute(text('SELECT pg_advisory_unlock(:namespace,:slot)'),{'namespace':limits.IMPORT_NAMESPACE,'slot':slot})
            owner.commit()
    assert c.post('/api/v1/imports/events',files={'file':('busy.csv',file_content('csv'))}).status_code==200


def test_identical_file_busy_is_conflict_not_bad_file(login):
    c=login();payload=file_content('csv');key=int.from_bytes(hashlib.sha256(('events_journal'+hashlib.sha256(payload).hexdigest()).encode()).digest()[:8],'big',signed=True)
    with engine.connect() as owner:
        owner.execute(text('SELECT pg_advisory_lock(:key)'),{'key':key});owner.commit()
        try:
            response=c.post('/api/v1/imports/events',files={'file':('same.csv',payload)})
            assert response.status_code==409 and response.json()['detail']['code']=='IMPORT_ALREADY_PROCESSING'
            assert 'retry-after' in response.headers
        finally:owner.execute(text('SELECT pg_advisory_unlock(:key)'),{'key':key});owner.commit()
    assert c.post('/api/v1/imports/events',files={'file':('same.csv',payload)}).json()['inserted_rows']==1
    assert c.post('/api/v1/imports/events',files={'file':('same.csv',payload)}).json()['already_imported']


@pytest.mark.parametrize('failure',[RuntimeError,KeyboardInterrupt])
def test_exception_releases_slot_and_file_locks(failure):
    with pytest.raises(failure):
        with limits.import_admission(9123456789):raise failure('injected operator interruption')
    with limits.import_admission(9123456789):pass
    assert limits.DEADLINE.get() is None
    with engine.connect() as db:
        assert db.scalar(text("SELECT count(*) FROM pg_locks WHERE locktype='advisory' AND database=(SELECT oid FROM pg_database WHERE datname=current_database())"))==0


def test_processing_deadline_rolls_back_and_returns_actionable_status(login,monkeypatch):
    @contextmanager
    def expired():
        token=limits.DEADLINE.set(time.monotonic()-1)
        try:yield
        finally:limits.DEADLINE.reset(token)
    monkeypatch.setattr(import_registry,'processing_budget',expired)
    response=login().post('/api/v1/imports/events',files={'file':('late.csv',file_content('csv'))})
    assert response.status_code==408
    with SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(EventsJournal))==0
        record=db.scalar(select(DataImport));assert record.status=='failed' and 'deadline' in record.error_message
    assert limits.DEADLINE.get() is None


def test_workbook_uncompressed_limit_checked_before_decoder(tmp_path,monkeypatch):
    path=tmp_path/'oversized.xlsx'
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:z.writestr('large.xml',b'x'*2048)
    monkeypatch.setattr(get_settings(),'max_xlsx_uncompressed_bytes',1024)
    monkeypatch.setattr(formats,'load_workbook',lambda *a,**k:pytest.fail('Workbook must be rejected before decode'))
    with pytest.raises(ValueError,match='uncompressed'):list(formats.file_rows(path))


def test_workbook_member_count_rejected_before_zipinfo_allocation(tmp_path,monkeypatch):
    path=tmp_path/'members.xlsx'
    with zipfile.ZipFile(path,'w') as z:
        for n in range(4097):z.writestr(str(n),'')
    monkeypatch.setattr(formats.zipfile,'ZipFile',lambda *a,**k:pytest.fail('Central directory must be checked before ZipFile'))
    with pytest.raises(ValueError,match='directory'):list(formats.file_rows(path))


def test_deadlock_retry_rolls_back_first_attempt_and_preserves_idempotency(login,monkeypatch):
    from sqlalchemy.exc import OperationalError
    original=import_registry.load_events_journal;attempts=[]
    class ForcedDeadlock(Exception):sqlstate='40P01'
    def collide(*args,**kwargs):
        result=original(*args,**kwargs);attempts.append(1)
        if len(attempts)==1:raise OperationalError('controlled conflict',{},ForcedDeadlock())
        return result
    monkeypatch.setattr(import_registry,'load_events_journal',collide)
    response=login().post('/api/v1/imports/events',files={'file':('overlap.csv',file_content('csv'))})
    assert response.status_code==200 and len(attempts)==2
    with SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(EventsJournal))==1
        assert db.scalar(select(func.count()).select_from(DataImport))==1
        assert db.scalar(select(DataImport)).status=='completed'


def test_exhausted_transaction_retry_is_backpressure_with_zero_partial_rows(login,monkeypatch):
    from sqlalchemy.exc import OperationalError
    original=import_registry.load_events_journal
    class ForcedDeadlock(Exception):sqlstate='40P01'
    def collide(*args,**kwargs):
        original(*args,**kwargs)
        raise OperationalError('controlled conflict',{},ForcedDeadlock())
    monkeypatch.setattr(import_registry,'load_events_journal',collide)
    response=login().post('/api/v1/imports/events',files={'file':('overlap.csv',file_content('csv'))})
    assert response.status_code==429 and response.headers['retry-after']=='3'
    with SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(EventsJournal))==0
        assert db.scalar(select(DataImport)).status=='failed'


def catalogue_file(rows,format):
    import io,csv
    if format=='csv':
        out=io.StringIO();writer=csv.DictWriter(out,fieldnames=list(rows[0]));writer.writeheader();writer.writerows(rows)
        return out.getvalue().encode('utf-8')
    from openpyxl import Workbook
    book=Workbook(write_only=True);sheet=book.create_sheet();sheet.append(list(rows[0]))
    for row in rows:sheet.append(list(row.values()))
    out=io.BytesIO();book.save(out);return out.getvalue()


@pytest.mark.parametrize('kind',['objects','channels'])
@pytest.mark.parametrize('format',['csv','xlsx'])
def test_catalogue_duplicate_ids_use_last_row_and_count_unique_upserts(login,kind,format):
    row={'ид_объект':'2','иерархия_уровень':'','родитель':'','вид_объекта':'Тест','диспетчерское_название_объекта':'Первое'} if kind=='objects' else {
        'ид_канала_данных':'2','ид_объект':'1','тип_датчика':'Дым','тип_инж_системы':'Пожарная',
        'тег_инженерной_системы':'ТЕСТ','название_датчика':'Первое'}
    name='диспетчерское_название_объекта' if kind=='objects' else 'название_датчика'
    payload=catalogue_file([row,{**row,name:'Последнее'},row,{**row,name:'Последнее'}],format)
    response=login().post('/api/v1/imports/events?kind='+kind,files={'file':('catalogue.'+format,payload)})
    assert response.status_code==200,response.text
    assert response.json()['processed_rows']==4 and response.json()['inserted_rows']==1 and response.json()['skipped_duplicates']==3
    with SessionLocal() as db:
        model=ObjectCatalogue if kind=='objects' else ChannelCatalogue
        assert getattr(db.get(model,2),name)=='Последнее'


def test_large_object_catalogue_avoids_single_statement_parameter_limit(login,monkeypatch):
    from app.ingestion import csv_loader
    original=csv_loader._upsert_object_batch;sizes=[]
    def observe(db,rows):sizes.append(len(rows));return original(db,rows)
    monkeypatch.setattr(csv_loader,'_upsert_object_batch',observe)
    rows=[{'ид_объект':str(10000+i),'иерархия_уровень':'','родитель':'','вид_объекта':'Тест',
        'диспетчерское_название_объекта':'Синтетический объект'} for i in range(15000)]
    response=login().post('/api/v1/imports/events?kind=objects',files={'file':('large.csv',catalogue_file(rows,'csv'))})
    assert response.status_code==200,response.text
    assert max(sizes)<=1000 and len(sizes)==15
    with SessionLocal() as db:assert db.scalar(select(func.count()).select_from(ObjectCatalogue))==15001


def test_batched_object_catalogue_keeps_whole_file_atomic(login):
    rows=[{'ид_объект':str(10000+i),'иерархия_уровень':'','родитель':'','вид_объекта':'Тест',
        'диспетчерское_название_объекта':'Синтетический объект'} for i in range(3001)]
    rows[-1]['родитель']='invalid'
    response=login().post('/api/v1/imports/events?kind=objects',files={'file':('atomic.csv',catalogue_file(rows,'csv'))})
    assert response.status_code==400
    with SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(ObjectCatalogue))==1
        assert db.scalar(select(DataImport)).status=='failed'
