import io
from datetime import datetime, timezone
from openpyxl import load_workbook

def test_report_export_and_formula_safety(login):
    client=login()
    assert client.post("/api/v1/requests",json={"object_id":1,"title":"=1+1"}).status_code==200
    day=datetime.now(timezone.utc).date().isoformat()
    base=f"/api/v1/reports/repairs?date_from={day}&date_to={day}"
    data=client.get(base).json()
    assert data['source']=='preventive_requests' and data['timezone']=='UTC'
    assert data['generated_at'] and data['row_count']==1
    assert data["rows"][0]["title"]=="=1+1"
    xlsx=client.get(base+"&format=xlsx")
    assert xlsx.status_code==200
    workbook=load_workbook(io.BytesIO(xlsx.content),data_only=False)
    assert workbook.active.cell(4,3).value=="'=1+1"
    assert workbook.active.cell(4,3).data_type=="s"
    pdf=client.get(base+"&format=pdf")
    assert pdf.status_code==200
    assert pdf.content.startswith(b"%PDF-")

def test_report_period_validation(login):
    client=login()
    assert client.get("/api/v1/reports/observations?date_from=2026-10-01&date_to=2026-09-01").status_code==422

def test_report_empty_and_large_are_explicit(login, monkeypatch):
    from app.api.v1 import reports
    client=login()
    base='/api/v1/reports/repairs?date_from=2026-01-01&date_to=2026-12-31'
    assert client.get(base).json()['row_count']==0
    for format in ('pdf','xlsx'):
        assert client.get(base+'&format='+format).status_code==200
    for n in range(3):
        assert client.post('/api/v1/requests',json={'object_id':1,'title':f'Test {n}'}).status_code==200
    monkeypatch.setattr(reports,'ROW_LIMIT',2)
    report=client.get(base).json()
    assert report['truncated'] and report['row_count']==2 and report['row_limit']==2
