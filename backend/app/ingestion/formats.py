"""Bounded CSV/XLSX/JSON/XML decoding with no XML entities or spreadsheet formula evaluation."""
import csv
import io
import json
import struct
import zipfile
from pathlib import Path
from defusedxml import ElementTree
from openpyxl import load_workbook
from app.core.config import get_settings
from app.ingestion.limits import check_import_budget


def validate_workbook_directory(path):
    # Bound ZipInfo allocation before ZipFile parses the central directory.
    # XLSX imports do not need multi-disk/ZIP64 archives with >4096 members.
    with path.open('rb') as stream:
        stream.seek(0,2);size=stream.tell();stream.seek(max(0,size-65557));tail=stream.read()
    offset=tail.rfind(b'PK\x05\x06')
    if offset<0 or len(tail)-offset<22:raise ValueError('Malformed workbook ZIP directory')
    _,disk,central_disk,on_disk,total,directory_size,directory_offset,_=struct.unpack('<4s4H2LH',tail[offset:offset+22])
    if (disk or central_disk or on_disk!=total or total>4096 or directory_size>4*1024*1024
        or directory_offset+directory_size>size):
        raise ValueError('Workbook ZIP directory exceeds safe import limits')

def decode_records(data: bytes, format: str) -> list[dict]:
    if len(data) > get_settings().max_upload_bytes:
        raise ValueError("Input exceeds the configured byte limit")
    if format == "json":
        value = json.loads(data)
        if not isinstance(value, list) or any(not isinstance(row, dict) for row in value):
            raise ValueError("Expected an array of records")
        rows = value
    elif format == "xml":
        try:
            root = ElementTree.fromstring(data, forbid_dtd=True, forbid_entities=True, forbid_external=True)
        except Exception as exc:
            raise ValueError("Unsafe or malformed XML") from exc
        if root.tag != "records":
            raise ValueError("Expected <records><record>...</record></records>")
        rows = []
        for record in root:
            if record.tag != "record" or any(len(field) for field in record):
                raise ValueError("XML must contain flat records")
            keys = [field.tag for field in record]
            if len(keys) != len(set(keys)):
                raise ValueError("Duplicate XML field")
            rows.append({field.tag: field.text for field in record})
    else:
        raise ValueError("Expected JSON or XML")
    if len(rows) > get_settings().max_import_rows:
        raise ValueError("Too many records")
    return rows

def file_rows(path: Path):
    maximum = get_settings().max_import_rows
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            if not reader.fieldnames or len(reader.fieldnames) != len(set(reader.fieldnames)):
                raise ValueError("Missing or duplicate CSV columns")
            for index, row in enumerate(reader, 1):
                check_import_budget()
                if index > maximum or None in row:
                    raise ValueError("Too many rows or malformed CSV record")
                yield row
    elif suffix == ".xlsx":
        validate_workbook_directory(path)
        with zipfile.ZipFile(path) as archive:
            if len(archive.infolist())>4096 or sum(item.file_size for item in archive.infolist()) > get_settings().max_xlsx_uncompressed_bytes:
                raise ValueError("Workbook exceeds the configured uncompressed size or member limit")
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            sheet = workbook.worksheets[0]
            rows = sheet.iter_rows()
            first = next(rows, None)
            if first is None:
                raise ValueError("Empty workbook")
            headers = [str(cell.value or "").strip() for cell in first]
            if not all(headers) or len(headers) != len(set(headers)):
                raise ValueError("Missing or duplicate XLSX headers")
            for index, row in enumerate(rows, 1):
                check_import_budget()
                if index > maximum:
                    raise ValueError("Too many rows")
                if any(cell.data_type == "f" for cell in row):
                    raise ValueError("Formula cells are not accepted")
                if all(cell.value is None for cell in row):
                    continue
                yield dict(zip(headers, [cell.value for cell in row]))
        finally:
            workbook.close()
    elif suffix in {".json", ".xml"}:
        yield from decode_records(path.read_bytes(), suffix[1:])
    else:
        raise ValueError("Supported formats: CSV, XLSX, JSON, XML")

def normalize_to_csv(path: Path, output: Path, columns: set[str]):
    with output.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=sorted(columns), extrasaction="ignore")
        writer.writeheader()
        for index, row in enumerate(file_rows(path), 1):
            check_import_budget()
            missing = columns - row.keys()
            if missing:
                raise ValueError(f"Row {index}: missing columns {', '.join(sorted(missing))}")
            writer.writerow({key: ("true" if value is True else "false" if value is False else value)
                             for key, value in row.items()})
